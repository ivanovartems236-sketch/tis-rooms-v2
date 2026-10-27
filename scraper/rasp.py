# © 2026 Артём Иванов. Все права защищены.
# Копирование, изменение и распространение без письменного разрешения автора запрещены (см. LICENSE).
"""Загрузка и разбор расписания ТИС (филиал ДГТУ) с https://stis.su/schedule."""
import datetime as dt
import io
import os
import re
import time
import urllib.error
import urllib.parse
import urllib.request
import warnings

warnings.filterwarnings("ignore")
import openpyxl

SITE = "https://stis.su/schedule"
SKIP = ("rasp_exam", "zaschita")  # экзамены и ВКР — другой формат
DAYS = ["ПОНЕДЕЛЬНИК", "ВТОРНИК", "СРЕДА", "ЧЕТВЕРГ", "ПЯТНИЦА", "СУББОТА", "ВОСКРЕСЕНЬЕ"]
KINDS = {"л": 1, "пр": 2, "лаб": 3}  # 0 — не указано
ROOM_RE = re.compile(r"^[А-ЯЁA-Z]{1,3}(-\d+[А-Яа-яA-Za-z]?)?$")
_FIO = r"[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?\s+[А-ЯЁ]\.\s?[А-ЯЁ]\."
ONE_LINE_RE = re.compile(rf"^(.+?)\s+({_FIO}(?:\s*/\s*{_FIO})*)(?:\s+({ROOM_RE.pattern[1:-1]}))?$")
SHEET_RE =re.compile(r"(\d{1,2})\.(\d{1,2})\s*-\s*(\d{1,2})\.(\d{1,2})")
UA = {"User-Agent": "Mozilla/5.0 (tis-rooms schedule bot)"}


RETRY_DELAY = 45  # с; один повтор при сбое сети (разовый DNS-сбой 26.10)


def http_get(url):
    # в именах файлов на stis.su бывают пробелы и кириллица — без кодирования urllib их не скачивает
    req = urllib.request.Request(urllib.parse.quote(url, safe=":/?&=%#,"), headers=UA)
    for attempt in (1, 2):
        try:
            with urllib.request.urlopen(req, timeout=40) as r:  # проверка HTTPS включена
                return r.read()
        except urllib.error.HTTPError as e:
            if e.code < 500 or attempt == 2:  # 404 и прочие 4xx повтором не лечатся
                raise
        except OSError:  # DNS, таймаут, обрыв соединения (URLError — тоже OSError)
            if attempt == 2:
                raise
        time.sleep(RETRY_DELAY)


def schedule_links():
    html = http_get(SITE).decode("utf-8", "ignore")
    links = []
    for href in re.findall(r'href="([^"]+\.xlsx?)"', html, re.I):
        url = urllib.parse.urljoin(SITE, href)
        if url not in links and not any(s in url for s in SKIP):
            links.append(url)
    return links


def download_all(folder, log=print):
    """Возвращает (скачано, не скачано). Ссылки с 404 — мёртвые на самом сайте, их не считаем ошибкой.
    На первой настоящей ошибке останавливаемся: сборка всё равно не пойдёт, а повторы съели бы timeout workflow."""
    os.makedirs(folder, exist_ok=True)
    links = schedule_links()
    ok, failed = 0, 0
    for i, url in enumerate(links):
        name = f"{i:02d}_" + re.sub(r'[\\/:*?"<>|,\s]+', "_", urllib.parse.unquote(url.split("/")[-1]))
        try:
            data = http_get(url)
            with open(os.path.join(folder, name), "wb") as fh:
                fh.write(data)
            ok += 1
        except urllib.error.HTTPError as e:
            log(f"ошибка {url}: {e}")
            if e.code != 404:
                failed += 1
                break
        except Exception as e:
            log(f"ошибка {url}: {e}")
            failed += 1
            break
    log(f"скачано {ok} из {len(links)}")
    return ok, failed


def sheet_dates(title, ref):
    """Диапазон дат листа «19.10-25.10» с учётом учебного года относительно ref."""
    m = SHEET_RE.search(title)
    if not m:
        return None
    d1, m1, d2, m2 = map(int, m.groups())
    y = ref.year
    if m1 >= 8 and ref.month < 7:
        y -= 1
    elif m1 < 8 and ref.month >= 8:
        y += 1
    try:
        s = dt.date(y, m1, d1)
        e = dt.date(y if m2 >= m1 else y + 1, m2, d2)
    except ValueError:
        return None
    return s, e


def clean_teacher(s):
    """Опечатки в ФИО из исходного Excel: «Черепов П.В..», «Сапилкина В.И .», «Петров О. П.» → один человек."""
    out = []
    for t in s.split("/"):
        t = re.sub(r"\s+", " ", t).strip()
        t = re.sub(r"\s+\.", ".", t)             # «В.И .» → «В.И.»
        t = re.sub(r"\.{2,}", ".", t)              # «П.В..» → «П.В.»
        t = re.sub(r"(?<=[А-ЯЁ]\.)\s+(?=[А-ЯЁ]\.)", "", t)  # «О. П.» → «О.П.»
        if t:
            out.append(t)
    return " / ".join(out)


def parse_cell(v):
    lines = [l.strip() for l in str(v).split("\n") if l.strip()]
    m = re.match(r"^(лаб|пр|л)\.\s*(.*)$", lines[0])
    kind, subj = (KINDS[m.group(1)], m.group(2)) if m else (0, lines[0])
    if len(lines) == 1:
        # заочка: «Предмет Фамилия И.О. К-204» одной строкой
        m = ONE_LINE_RE.match(subj)
        if m:
            return kind, m.group(1).strip(), clean_teacher(m.group(2)), m.group(3) or ""
    toks = (lines[-1] if len(lines) > 1 else "").split()
    room = toks.pop() if toks and ROOM_RE.match(toks[-1]) else ""
    if len(lines) > 2:
        subj += " " + " ".join(lines[1:-1])
    teacher = clean_teacher(" ".join(toks))
    if teacher.startswith("("):  # «(получение паролей от ЭИОС…)» — пояснение к занятию, а не преподаватель
        subj, teacher = subj + " " + teacher, ""
    return kind, re.sub(r"\s+", " ", subj).strip(), teacher, room


def parse_sheet(ws):
    """[(weekday, pair, time, group, kind, subj, teacher, room)]"""
    hdr = next((r for r in range(1, 30) if str(ws.cell(r, 2).value or "").strip() == "Пара"), None)
    if not hdr:
        return []
    groups = {c: str(ws.cell(hdr, c).value).strip()
              for c in range(4, ws.max_column + 1) if str(ws.cell(hdr, c).value or "").strip()}
    merged = {(mr.min_row, mr.min_col): range(mr.min_col, mr.max_col + 1)
              for mr in ws.merged_cells.ranges if mr.min_row > hdr and mr.max_col > mr.min_col}
    out, day = [], None
    for r in range(hdr + 1, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "").upper().split()  # бывает «05 октября\nПОНЕДЕЛЬНИК»
        found = [w for w in a if w in DAYS]
        if found:
            day = DAYS.index(found[-1])
        try:
            pair = int(str(ws.cell(r, 2).value).strip())
        except (TypeError, ValueError):
            continue
        if day is None:
            continue
        time = re.sub(r"\s+", "", str(ws.cell(r, 3).value or ""))
        for c in groups:
            v = ws.cell(r, c).value
            if v is None or not str(v).strip():
                continue
            kind, subj, teacher, room = parse_cell(v)
            for cc in merged.get((r, c), [c]):
                if cc in groups:
                    out.append((day, pair, time, groups[cc], kind, subj, teacher, room))
    return out


def load_book(path):
    with open(path, "rb") as fh:
        data = fh.read()
    return openpyxl.load_workbook(io.BytesIO(data))  # файлы .xls на сайте на самом деле xlsx


def lessons_from_folder(folder, ref, keep_days=10, log=print):
    """Плоский список занятий с конкретными датами: (date, pair, time, group, kind, subj, teacher, room)."""
    seen = set()
    for f in sorted(os.listdir(folder)):
        try:
            wb = load_book(os.path.join(folder, f))
        except Exception as e:
            log(f"пропуск {f}: {e}")
            continue
        for ws in wb.worksheets:
            rng = sheet_dates(ws.title, ref)
            if not rng or rng[1] < ref - dt.timedelta(days=keep_days):
                continue
            s, e = rng
            days = {}
            d = s
            while d <= e:
                days[d.weekday()] = d
                d += dt.timedelta(days=1)
            for wd, pair, time, *rest in parse_sheet(ws):
                if wd in days:
                    seen.add((days[wd].isoformat(), pair, time, *rest))
    return sorted(seen)
